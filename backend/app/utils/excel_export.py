"""Excel workbook generation with openpyxl."""

import os
from datetime import date, datetime
from typing import Any, Dict, List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=10)
BODY_FONT = Font(size=10)
THIN = Side(style="thin", color="D1D5DB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

MAX_COLUMN_WIDTH = 45


def _stringify(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ", timespec="seconds") if isinstance(
            value, datetime
        ) else value.isoformat()
    if isinstance(value, (dict, list)):
        return str(value)
    if hasattr(value, "value"):  # enum
        return value.value
    return value


def write_sheet(
    workbook: Workbook,
    title: str,
    rows: List[Dict[str, Any]],
    columns: Optional[List[str]] = None,
    first: bool = False,
) -> None:
    sheet = workbook.active if first else workbook.create_sheet()
    sheet.title = title[:31]  # Excel's hard limit on sheet names

    if not rows:
        sheet["A1"] = "No data for this selection."
        sheet["A1"].font = BODY_FONT
        return

    headers = columns or list(rows[0].keys())
    for index, header in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=index, value=header.replace("_", " ").title())
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")

    for row_index, row in enumerate(rows, start=2):
        for column_index, header in enumerate(headers, start=1):
            cell = sheet.cell(
                row=row_index, column=column_index, value=_stringify(row.get(header))
            )
            cell.font = BODY_FONT
            cell.border = BORDER

    for index, header in enumerate(headers, start=1):
        longest = max(
            [len(str(header))]
            + [len(str(_stringify(row.get(header)) or "")) for row in rows[:200]]
        )
        sheet.column_dimensions[get_column_letter(index)].width = min(
            longest + 3, MAX_COLUMN_WIDTH
        )

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = (
        f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    )


def build_workbook(
    sheets: Dict[str, List[Dict[str, Any]]], output_path: str
) -> str:
    os.makedirs(os.path.dirname(output_path) or ".", exist_ok=True)
    workbook = Workbook()

    for index, (title, rows) in enumerate(sheets.items()):
        write_sheet(workbook, title, rows, first=(index == 0))

    workbook.save(output_path)
    return output_path
